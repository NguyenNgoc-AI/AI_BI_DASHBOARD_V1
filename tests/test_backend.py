import json
import io
import sqlite3
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from pptx import Presentation

from backend.auth import authenticate_user, create_access_token, create_user, decode_access_token
from backend.chatbot.llm_router import DeterministicAdapter, LLMRouter
from backend.config import settings
from backend.chatbot.text_to_sql import execute_read_only, question_to_sql, validate_select_sql
from backend.dashboard_engine.chart_generator import revenue_by_date
from backend.dashboard_engine.data_pipeline import ingest
from backend.dashboard_engine.kpi_calculator import calculate_kpis
from backend.dashboard_engine.report_exporter import export_report
from backend.eda.profiling import profile_dataset
from backend.eda.quality_checker import check_quality
from backend.schema_learning.rule_parser import evaluate_rules
from backend.schema_learning.schema_extractor import infer_database_schema, infer_schema, pearson_correlation
from backend.synthetic_generator.pipeline import generate_synthetic_dataset
from backend.synthetic_generator.generator_model import create_generator


CSV = b"date,order_id,product,revenue,cost\n2026-10-01,O001,Bag A,100,60\n2026-10-01,O002,Bag B,200,80\n2026-10-02,O003,Bag A,150,70\n"


class BackendTests(unittest.TestCase):
    def setUp(self):
        self.rows = ingest(CSV, "sales.csv")

    def test_end_to_end_analytics(self):
        self.assertEqual(calculate_kpis(self.rows)["net_profit"], 240)
        self.assertEqual(revenue_by_date(self.rows)[0]["revenue"], 300)
        self.assertEqual(profile_dataset(self.rows)["row_count"], 3)
        self.assertEqual(check_quality(self.rows)["score"], 1.0)
        self.assertEqual(infer_schema(self.rows)["columns"]["revenue"]["type"], "float")

    def test_ingest_json_and_reject_bad_data(self):
        content = json.dumps(self.rows).encode()
        self.assertEqual(len(ingest(content, "sales.json")), 3)
        with self.assertRaises(ValueError):
            ingest(b"date,order_id\n2026-01-01,1\n", "bad.csv")
        with self.assertRaises(ValueError):
            ingest(b"{}", "bad.json")
        with self.assertRaises(ValueError):
            ingest(b"{", "malformed.json")
        with self.assertRaises(ValueError):
            ingest(b"date,order_id,product,revenue,cost\n2026-01-01,1,A,nan,1\n", "bad.csv")

    def test_rules_and_synthetic_are_deterministic(self):
        rules = [{"field": "revenue", "operator": ">=", "value": 0}]
        self.assertEqual(evaluate_rules(self.rows, rules)["rate"], 1.0)
        first = generate_synthetic_dataset(self.rows, 3, seed=42)
        second = generate_synthetic_dataset(self.rows, 3, seed=42)
        self.assertEqual(first, second)
        self.assertEqual(set(first["validation"]), {"validity", "fidelity", "privacy", "utility", "details"})
        self.assertNotEqual(first["rows"][0]["order_id"], self.rows[0]["order_id"])
        self.assertEqual(first["model"], "bootstrap")
        self.assertEqual(create_generator("ctgan").model, "ctgan")
        with self.assertRaises(ValueError):
            create_generator("unknown")

        evaluated = generate_synthetic_dataset(self.rows, 20, seed=42, target_column="revenue")
        self.assertEqual(evaluated["validation"]["details"]["utility_method"], "tstr_regression")

    def test_auth_hashes_password_and_signs_token(self):
        with tempfile.TemporaryDirectory() as directory:
            database = str(Path(directory) / "auth.db")
            user = create_user("ngoc_test", "strong-password", database)
            self.assertIsNone(authenticate_user("ngoc_test", "wrong-password", database))
            self.assertEqual(authenticate_user("ngoc_test", "strong-password", database)["role"], "user")
            token = create_access_token(user)
            self.assertEqual(decode_access_token(token)["username"], "ngoc_test")
            with self.assertRaises(ValueError):
                decode_access_token(token[:-1] + ("A" if token[-1] != "A" else "B"))
            with self.assertRaises(ValueError):
                decode_access_token(create_access_token(user, expires_minutes=-1))

    def test_text_to_sql_is_read_only(self):
        columns = {"order_id", "product", "revenue", "cost"}
        sql = question_to_sql("Doanh thu theo sản phẩm", {})
        validated = validate_select_sql(sql, {"sales"}, columns)
        self.assertIn("LIMIT 500", validated)
        for unsafe in ("DROP TABLE sales", "SELECT * FROM unknown", "SELECT secret FROM sales", "SELECT * FROM sales; DELETE FROM sales"):
            with self.assertRaises(ValueError):
                validate_select_sql(unsafe, {"sales"}, columns)

        connection = sqlite3.connect(":memory:")
        connection.execute("CREATE TABLE sales(revenue REAL)")
        connection.execute("INSERT INTO sales VALUES (100)")
        result = execute_read_only(connection, "SELECT SUM(revenue) AS total_revenue FROM sales LIMIT 500")
        self.assertEqual(result[0]["total_revenue"], 100)
        with self.assertRaises(sqlite3.DatabaseError):
            connection.execute("DELETE FROM sales")

        class SQLAdapter:
            def generate(self, *_args, **_kwargs):
                return "```sql\nSELECT SUM(revenue) AS total_revenue FROM sales\n```"

        generated = question_to_sql("Tổng doanh thu", {"columns": {}}, SQLAdapter())
        self.assertEqual(generated, "SELECT SUM(revenue) AS total_revenue FROM sales")

    def test_offline_llm_and_reports(self):
        kpis = calculate_kpis(self.rows)
        answer = DeterministicAdapter().generate("system", "question", {"kpis": kpis})
        self.assertIn("450.0", answer)
        report = {"kpis": {"product": "=1+1", **kpis}}
        csv_content, _ = export_report(report, "csv")
        self.assertIn(b"'=1+1", csv_content)
        for format in ("json", "html"):
            self.assertTrue(export_report(report, format)[0])
        pdf, _ = export_report(report, "pdf")
        self.assertTrue(pdf.startswith(b"%PDF"))
        xlsx, _ = export_report(report, "xlsx")
        self.assertEqual(load_workbook(io.BytesIO(xlsx))["KPI"]["A1"].value, "Metric")
        pptx, _ = export_report(report, "pptx")
        self.assertEqual(Presentation(io.BytesIO(pptx)).slides[0].shapes.title.text, "AI BI Report")
        with self.assertRaises(ValueError):
            export_report(report, "docx")

    @patch("backend.chatbot.llm_router.httpx.post")
    def test_openai_router_uses_responses_api(self, post):
        response = Mock()
        response.json.return_value = {
            "output": [{"type": "message", "content": [{"type": "output_text", "text": "Phân tích thật"}]}]
        }
        post.return_value = response
        router = LLMRouter(replace(settings, llm_provider="openai", openai_api_key="test-key"))
        answer = router.generate("system", "question", {"kpis": {}})
        self.assertEqual(answer, "Phân tích thật")
        self.assertEqual(post.call_args.args[0], "https://api.openai.com/v1/responses")
        self.assertEqual(router.used_provider, "openai")

    def test_hardening_boundaries(self):
        self.assertIsNone(pearson_correlation([1.0, 1.0], [2.0, 3.0]))
        report = {"kpis": {"space": "  =1+1", "tab": "\t@SUM(A1:A2)"}}
        csv_content, _ = export_report(report, "csv")
        self.assertIn(b"'  =1+1", csv_content)
        self.assertIn(b"'\t@SUM", csv_content)

    def test_multi_table_relationship_inference(self):
        schema = infer_database_schema(
            {
                "customers": [{"customer_id": "C1"}, {"customer_id": "C2"}],
                "orders": [{"order_id": "O1", "customer_id": "C1"}, {"order_id": "O2", "customer_id": "C2"}],
            }
        )
        self.assertIn("customer_id", schema["primary_key_candidates"]["customers"])
        self.assertEqual(schema["foreign_keys"][0]["child_table"], "orders")

    def test_llm_agent_generator_validates_json_and_replaces_ids(self):
        class FakeLLM:
            def generate(self, *_args, **_kwargs):
                return json.dumps([self_rows[0]])

        self_rows = self.rows
        result = generate_synthetic_dataset(self.rows, 1, model="llm_agent", llm=FakeLLM())
        self.assertEqual(result["model"], "llm_agent")
        self.assertEqual(result["rows"][0]["order_id"], "synthetic-order_id-1")

    @patch("backend.chatbot.llm_router.httpx.post")
    def test_gemini_and_llama_provider_contracts(self, post):
        gemini_response = Mock()
        gemini_response.json.return_value = {"candidates": [{"content": {"parts": [{"text": "Gemini OK"}]}}]}
        post.return_value = gemini_response
        gemini = LLMRouter(replace(settings, llm_provider="gemini", gemini_api_key="test-key"))
        self.assertEqual(gemini.generate("system", "question"), "Gemini OK")
        self.assertIn("generativelanguage.googleapis.com", post.call_args.args[0])

        llama_response = Mock()
        llama_response.json.return_value = {"choices": [{"message": {"content": "Llama OK"}}]}
        post.return_value = llama_response
        llama = LLMRouter(replace(settings, llm_provider="llama", llama_base_url="http://localhost:11434"))
        self.assertEqual(llama.generate("system", "question"), "Llama OK")
        self.assertEqual(post.call_args.args[0], "http://localhost:11434/v1/chat/completions")


if __name__ == "__main__":
    unittest.main()
