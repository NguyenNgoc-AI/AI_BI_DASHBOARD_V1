"""Export dashboard reports to formats commonly used by BI teams."""

import csv
from html import escape
import io
import json


def _metric_rows(report: dict) -> list[tuple[str, object]]:
    return list(report.get("kpis", {}).items())


def build_report_payload(profile: dict, quality: dict, kpis: dict, charts: list[dict]) -> dict:
    return {"profile": profile, "quality": quality, "kpis": kpis, "charts": charts}


def export_json(report: dict) -> bytes:
    return json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8")


def _safe_csv_cell(value: object) -> object:
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_csv(report: dict) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["metric", "value"])
    for key, value in report.get("kpis", {}).items():
        writer.writerow([_safe_csv_cell(key), _safe_csv_cell(value)])
    return output.getvalue().encode("utf-8-sig")


def export_html(report: dict) -> bytes:
    rows = "".join(f"<tr><th>{escape(str(key))}</th><td>{escape(str(value))}</td></tr>" for key, value in _metric_rows(report))
    return f"<!doctype html><meta charset='utf-8'><title>AI BI Report</title><h1>AI BI Report</h1><table>{rows}</table>".encode()


def export_pdf(report: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas

    output = io.BytesIO()
    canvas = Canvas(output, pagesize=A4)
    _, height = A4
    y = height - 50
    canvas.setTitle("AI BI Report")
    canvas.setFont("Helvetica-Bold", 18)
    canvas.drawString(50, y, "AI BI Report")
    canvas.setFont("Helvetica", 11)
    for key, value in _metric_rows(report):
        y -= 24
        if y < 50:
            canvas.showPage()
            canvas.setFont("Helvetica", 11)
            y = height - 50
        canvas.drawString(50, y, f"{key}: {value}")
    canvas.save()
    return output.getvalue()


def export_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KPI"
    sheet.append(["Metric", "Value"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for key, value in _metric_rows(report):
        sheet.append([_safe_csv_cell(key), _safe_csv_cell(value)])
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 22
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_pptx(report: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "AI BI Report"
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(5.2))
    frame = box.text_frame
    frame.clear()
    for index, (key, value) in enumerate(_metric_rows(report)):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"{key}: {value}"
        paragraph.font.size = Pt(18)
        paragraph.space_after = Pt(10)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def export_report(report: dict, format: str) -> tuple[bytes, str]:
    exporters = {
        "json": (export_json, "application/json"),
        "csv": (export_csv, "text/csv; charset=utf-8"),
        "html": (export_html, "text/html; charset=utf-8"),
        "pdf": (export_pdf, "application/pdf"),
        "xlsx": (export_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "pptx": (export_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    }
    if format not in exporters:
        raise ValueError(f"Unsupported report format: {format}")
    exporter, media_type = exporters[format]
    return exporter(report), media_type
